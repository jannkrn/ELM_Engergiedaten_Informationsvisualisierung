from __future__ import annotations

from datetime import datetime
from pathlib import Path
import re
import tempfile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "report" / "final"
OUTPUT = SOURCE / "Projektbericht_Jann_Koerner_bearbeitbar.docx"
FIGURES = ROOT / "experiments" / "figures" / "final"

BLACK = "000000"
NAVY = "1F334D"
WHITE = "FFFFFF"
LIGHT_BLUE = "EEF3F8"
LIGHT_GRAY = "F5F5F5"
BORDER = "D9D9D9"
MUTED = "5B6472"

CITATIONS = {
    "aigner2011time": 1,
    "smardFlows": 2,
    "elmArchitecture": 3,
    "energychartsOpenapi": 4,
    "jaeger2016outage": 5,
    "keim2000pixel": 6,
    "munzner2009nested": 7,
    "north2000snap": 8,
    "postgrestPagination": 9,
    "steiger2013smartgrid": 10,
    "stoffel2012amplio": 11,
}

REFS = {
    "sec:einleitung": "Abschnitt 1",
    "sec:hintergrund": "Abschnitt 1.1",
    "sec:zielgruppen": "Abschnitt 1.2",
    "sec:beitraege": "Abschnitt 1.3",
    "sec:daten": "Abschnitt 2",
    "sec:visualisierungen": "Abschnitt 3",
    "sec:aufgaben": "Abschnitt 3.1",
    "sec:anforderungen": "Abschnitt 3.2",
    "sec:praesentation": "Abschnitt 3.3",
    "sec:vis1": "Abschnitt 3.3.1",
    "sec:vis2": "Abschnitt 3.3.2",
    "sec:vis3": "Abschnitt 3.3.3",
    "sec:interaktion": "Abschnitt 3.4",
    "sec:implementierung": "Abschnitt 4",
    "sec:anwendungsfaelle": "Abschnitt 5",
    "sec:case1": "Abschnitt 5.1",
    "sec:case2": "Abschnitt 5.2",
    "sec:case3": "Abschnitt 5.3",
    "sec:related": "Abschnitt 6",
    "sec:fazit": "Abschnitt 7",
    "fig:overview": "Abbildung 1",
    "fig:architecture": "Abbildung 2",
    "fig:case-france": "Abbildung 3",
    "fig:case-price": "Abbildung 4",
    "fig:case-matrix": "Abbildung 5",
    "tab:views": "Tabelle 1",
    "tab:generation": "Tabelle 2",
    "tab:tasks": "Tabelle 3",
}

FIGURE_CAPTIONS = {
    "fig:overview": (
        "Abbildung 1: Gesamtansicht der Elm-Anwendung mit Netzwerk, Zeitreihe und Pixelmatrix. "
        "Alle Zahlen und SVG-Geometrien werden von Elm aus dem HTTP-Datensatz erzeugt."
    ),
    "fig:architecture": "Abbildung 2: Modulstruktur und gerichteter Datenfluss der Elm-Anwendung.",
    "fig:case-france": (
        "Abbildung 3: Frankreich ist ausgewählt. Knoten und Kante bleiben hervorgehoben; "
        "die rote Richtung steht für physischen Import nach Deutschland."
    ),
    "fig:case-price": (
        "Abbildung 4: Monatliches Preisminimum. Die gemeinsame Zeitmarke verbindet "
        "erneuerbare Erzeugung, Frankreich-Fluss und Detailwerte."
    ),
    "fig:case-matrix": (
        "Abbildung 5: Auswahl der Zelle Dänemark, 03.05.2025 12 Uhr UTC. "
        "Das goldene Fadenkreuz macht beide gewählten Dimensionen sichtbar."
    ),
}

REFERENCES = [
    "[1] Wolfgang Aigner, Silvia Miksch, Heidrun Schumann und Christian Tominski. "
    "Visualization of Time-Oriented Data. Springer, 2011.",
    "[2] Bundesnetzagentur. Stromhandel und physikalischer Stromfluss. "
    "https://www.smard.de/page/home/wiki-article/518/596/stromhandel-und-physikalischer-stromfluss, "
    "2026. Abgerufen am 10.09.2026.",
    "[3] Elm Project. The Elm Architecture. https://guide.elm-lang.org/architecture/. "
    "Abgerufen am 10.09.2026.",
    "[4] Fraunhofer ISE. Energy-Charts API: OpenAPI Description. "
    "https://api.energy-charts.info/openapi.json, 2026. Abgerufen am 10.09.2026.",
    "[5] Alexander Jäger, Sebastian Mittelstädt, Daniela Oelke, Sonja Sander, Axel Platz, "
    "Gies Bouwman und Daniel A. Keim. Lessons on combining topology and geography - visual "
    "analytics for electrical outage management. EuroVis Workshop on Visual Analytics, 2016.",
    "[6] Daniel A. Keim. Designing pixel-oriented visualization techniques: Theory and applications. "
    "IEEE Transactions on Visualization and Computer Graphics, 6(1):59-78, 2000.",
    "[7] Tamara Munzner. A nested process model for visualization design and validation. "
    "IEEE Transactions on Visualization and Computer Graphics, 15(6):921-928, 2009.",
    "[8] Chris North und Ben Shneiderman. Snap-together visualization: can users construct and "
    "operate coordinated visualizations? International Journal of Human-Computer Studies, "
    "53(5):715-739, 2000.",
    "[9] PostgREST Project. Tables and views: Limits and pagination. "
    "https://docs.postgrest.org/en/stable/references/api/tables_views.html. "
    "Abgerufen am 10.09.2026.",
    "[10] Michael Steiger, Thorsten May, James Davey und Jörn Kohlhammer. Visual analysis of "
    "expert systems for smart grid monitoring. EuroVis Workshop on Visual Analytics, "
    "Seiten 43-47, 2013.",
    "[11] Andreas Stoffel, Lei Zhang, Stefan H. Weber und Daniel A. Keim. Amplio VQA - a web "
    "based visual query analysis system for micro grid energy mix planning. EuroVis Workshop "
    "on Visual Analytics, Seiten 73-77, 2012.",
]


def set_font(run, size=None, bold=None, italic=None, color=BLACK, name="Calibri"):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), name)
    rpr.rFonts.set(qn("w:hAnsi"), name)
    rpr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color=BORDER, size="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tag = OxmlElement("w:tblHeader")
    tag.set(qn("w:val"), "true")
    tr_pr.append(tag)


def set_table_widths(table, widths_in):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_in:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(round(width * 1440)))
        grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(round(widths_in[i] * 1440)))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            set_cell_border(cell)


def configure_document(doc):
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.35)
    sec.bottom_margin = Cm(2.1)
    sec.left_margin = Cm(2.45)
    sec.right_margin = Cm(2.45)
    sec.footer_distance = Cm(1.0)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_after = Pt(5.5)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.widow_control = True

    title = doc.styles["Title"]
    title.font.name = "Calibri"
    title.font.size = Pt(26)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title.paragraph_format.space_after = Pt(14)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    heading_specs = {
        "Heading 1": (16, 16, 8),
        "Heading 2": (13, 12, 5),
        "Heading 3": (11.5, 9, 4),
    }
    for name, (size, before, after) in heading_specs.items():
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(10.5)
        style.paragraph_format.left_indent = Cm(0.65)
        style.paragraph_format.first_line_indent = Cm(-0.35)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = 1.05

    for level in (1, 2, 3):
        name = f"TOC {level}"
        if name not in [style.name for style in doc.styles]:
            style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        else:
            style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(8.6 if level == 1 else 8.2)
        style.font.bold = level == 1
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.line_spacing = 1.0

    caption = doc.styles["Caption"]
    caption.font.name = "Calibri"
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(MUTED)
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.keep_with_next = False

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(footer, "PAGE", "1")
    sec.different_first_page_header_footer = True
    sec.first_page_footer.paragraphs[0].text = ""
    doc.settings.update_fields_on_open = True


def add_field(paragraph, instruction, result=""):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = result
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_font(run, size=9)


def add_hyperlink(paragraph, text, url):
    rel = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    style = OxmlElement("w:rStyle")
    style.set(qn("w:val"), "Hyperlink")
    rpr.append(style)
    run.append(rpr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def clean_tex(text):
    text = text.replace("{\\ss}", "ß")
    text = text.replace("``", "„").replace("''", "“")
    text = text.replace("---", "-").replace("--", "-")
    text = text.replace("~", " ").replace("\\,", "")
    text = text.replace("\\%", "%").replace("\\_", "_").replace("\\&", "&")
    text = text.replace("\\textbackslash", "\\")
    text = re.sub(r"\\projectlink\{([^{}]*)\}", r"\1", text)
    text = re.sub(r"\\href\{[^{}]*\}\{([^{}]*)\}", r"\1", text)
    text = re.sub(r"\\url\{([^{}]*)\}", r"\1", text)

    def citation(match):
        nums = [CITATIONS[k.strip()] for k in match.group(1).split(",") if k.strip() in CITATIONS]
        return "[" + ", ".join(str(n) for n in nums) + "]"

    def crossref(match):
        values = [REFS.get(k.strip(), k.strip()) for k in match.group(1).split(",")]
        if len(values) == 1:
            return values[0]
        prefix = "Abbildungen" if all(v.startswith("Abbildung") for v in values) else "Abschnitte"
        return prefix + " " + ", ".join(v.split(" ", 1)[1] for v in values)

    text = re.sub(r"\\cite\{([^{}]+)\}", citation, text)
    text = re.sub(r"\\[cC]ref\{([^{}]+)\}", crossref, text)
    for macro in ("code", "emph", "textit", "textbf", "sffamily"):
        text = re.sub(rf"\\{macro}\{{([^{{}}]*)\}}", r"\1", text)
    text = text.replace("\\gw", " GW").replace("\\mwh", " EUR/MWh")
    text = text.replace("\\times", " × ").replace("\\min", "min")
    text = text.replace("\\quad", " ").replace("\\sum", "Summe")
    text = re.sub(r"\\text\{([^{}]*)\}", r"\1", text)
    text = re.sub(r"\\underbrace\{([^{}]*)\}_\{[^{}]*\}", r"\1", text)
    text = text.replace("\\left", "").replace("\\right", "")
    text = text.replace("\\(", "").replace("\\)", "")
    text = text.replace("$", "")
    text = text.replace("{,}", ",")
    text = re.sub(r"_\{([^{}]*)\}", r"_\1", text)
    text = re.sub(r"\^\{([^{}]*)\}", r"^\1", text)
    text = text.replace("\\", "")
    text = text.replace("{", "").replace("}", "")
    text = re.sub(r"\s+", " ", text).strip()
    return text


def add_body_paragraph(doc, text, *, italic=False, centered=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(clean_tex(text))
    set_font(run, size=10.5, italic=italic)
    return p


def add_heading(doc, text, level):
    p = doc.add_paragraph(style=f"Heading {level}")
    run = p.add_run(clean_tex(text))
    set_font(run, bold=True, size={1: 16, 2: 13, 3: 11.5}[level])
    return p


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    repeat_header(table.rows[0])
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(header)
        set_font(run, size=9, bold=True, color=WHITE)
    for r_idx, values in enumerate(rows):
        cells = table.add_row().cells
        if r_idx % 2:
            for cell in cells:
                set_cell_shading(cell, LIGHT_BLUE)
        for i, value in enumerate(values):
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(clean_tex(value))
            set_font(run, size=8.6)
    set_table_widths(table, widths)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)
    return table


def add_caption(doc, text):
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_font(run, size=9, italic=True, color=MUTED)
    return p


def add_figure(doc, image_path, caption, width_cm=15.5, alt_text=""):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(image_path), width=Cm(width_cm))
    if alt_text:
        shape._inline.docPr.set("descr", alt_text)
    add_caption(doc, caption)


def add_code_block(doc, code):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_GRAY)
    set_cell_border(cell, BORDER)
    set_cell_margins(cell, 100, 130, 100, 130)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_together = True
    run = p.add_run(code.rstrip())
    set_font(run, size=8.5, name="Consolas")
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def crop_report_page(page_number, box, target):
    source = SOURCE / "rendered" / f"page-{page_number:02d}.png"
    with Image.open(source) as im:
        left, top, right, bottom = box
        crop = im.crop((int(im.width * left), int(im.height * top), int(im.width * right), int(im.height * bottom)))
        crop.save(target)


def insert_special_table(doc, label):
    if label == "tab:views":
        doc.add_page_break()
        add_caption(doc, "Tabelle 1: Verwendete PostgreSQL-Views und semantische Einheiten.")
        add_table(doc, ["View", "Inhalt", "Einheit"], [
            ["v_cbpf", "bilateraler physischer Grenzfluss; positiv Import nach Deutschland, negativ Export aus Deutschland", "GW"],
            ["v_cbet", "bilateraler grenzüberschreitender Handel mit derselben Vorzeichenkonvention", "GW"],
            ["v_price", "Day-Ahead-Preis der Gebotszone DE-LU", "EUR/MWh"],
            ["v_totalpower", "gesamte deutsche Nettoerzeugung einschließlich industrieller Eigenerzeugung, getrennt nach Technologien", "MW"],
        ], [1.15, 4.45, 0.9])
    elif label == "tab:generation":
        doc.add_page_break()
        add_caption(doc, "Tabelle 2: Gruppierung der Erzeugungstechnologien.")
        add_table(doc, ["Gruppe", "Enthaltene Technologien"], [
            ["Erneuerbare", "Wind an Land und auf See, Solar, Biomasse, Laufwasser, Speicherwasser und Geothermie"],
            ["Kohle", "Braunkohle und Steinkohle"],
            ["Gas", "fossiles Gas"],
            ["Sonstige", "Öl, Kohlegase, Abfall, Pumpspeicher, Kernenergie und sonstige Erzeugung"],
        ], [1.45, 5.05])
    elif label == "tab:tasks":
        doc.add_page_break()
        add_caption(doc, "Tabelle 3: Anwendungsaufgaben und dafür benötigte mentale Modelle.")
        add_table(doc, ["ID", "Aufgabe", "Aufzubauendes mentales Modell"], [
            ["A1", "Für eine Stunde erkennen, zwischen welchen Partnerländern und Deutschland physische Leistung in welcher Richtung und Größenordnung fließt.", "Deutschland als Zentrum eines gerichteten, gewichteten Netzes; Länderbeziehungen statt isolierter Tabellenzeilen."],
            ["A2", "Zeitliche Muster und Ausnahmen in Erzeugungsmix, Preis und dem Fluss eines ausgewählten Länderpaars untersuchen.", "Gemeinsame Zeitachse mit mehreren synchronen quantitativen Größen und klar getrennten Skalen."],
            ["A3", "In vielen Länder-Stunden-Kombinationen Perioden mit ähnlicher Richtung, starken Beträgen oder Richtungswechseln finden und als Detailzustand auswählen.", "Dichte zweidimensionale Übersicht mit Land und Zeit als orthogonalen Dimensionen."],
        ], [0.4, 3.0, 3.1])
    elif label == "app:feedback":
        add_table(doc, ["Rückmeldung", "Umsetzung"], [
            ["Text ausbauen; Text direkt unter Einleitung", "Zielproblem und Forschungsfrage stehen vor Unterabschnitt 1.1; alle Kapitel wurden fachlich ausgearbeitet."],
            ["Forschungsfrage gehört in Kapitel 1", "In Abschnitt 1 als hervorgehobene Leitfrage platziert."],
            ["Anwendungshintergrund ergänzen", "Physischer Fluss, Handel, Gebotszone, Erzeugungsmix und Preis werden in Abschnitt 1.1 erläutert."],
            ["Einheiten nicht raten", "Offizielle OpenAPI als Quelle; MW/GW-Transformation, Vorzeichen und EUR/MWh in Abschnitt 2 dokumentiert."],
            ["Zeitraum begründen", "Vollständiger Mai 2025, Umfang und Kontraste in Abschnitt 2 begründet."],
            ["Anforderungen unabhängig von Techniken", "Erst Aufgaben und Anforderungen in Abschnitten 3.1 und 3.2, danach Visualisierungsauswahl."],
            ["Pfeilspitzen kleiner", "SVG-Marker von 6 auf 4 Einheiten reduziert."],
            ["Maximale Linienbreite größer", "Kodierung auf bis zu 21,5 SVG-Einheiten erweitert."],
            ["Plus/Minus in Zeitreihe erklären", "Sichtbare Vorzeichenlegende im Fluss-Teilplot und Erläuterung in Abschnitt 3.3.2."],
            ["Andere Zeiträume auswählbar?", "48 Stunden, 7 Tage, Monat sowie Vor-/Zurück-Navigation umgesetzt."],
            ["Weiße Matrixlücken prüfen", "Zellkonturen entfernt; Nullwerte hellgrau statt weiß kodiert."],
            ["Nutzen der Auswahl erklären", "Matrix wählt Land und Stunde gleichzeitig; goldene Zeilen-/Spaltenmarkierung in Abschnitt 3.4."],
            ["Farbtöne zeilenübergreifend vergleichbar?", "Ein globales Monatsmaximum steuert alle Matrixzellen; numerische Legende ergänzt."],
            ["Abbildung für jeden Anwendungsfall", "Annotierte Abbildungen 3 bis 5 ergänzt."],
        ], [2.25, 4.25])


def insert_special_figure(doc, label, temp_dir):
    if label == "fig:overview":
        add_figure(doc, FIGURES / "01_uebersicht.png", FIGURE_CAPTIONS[label], 15.7, "Gesamtansicht der interaktiven Elm-Anwendung")
        return
    crops = {
        "fig:architecture": (12, (0.10, 0.265, 0.90, 0.515), 13.8, "Modul- und Datenflussdiagramm der Elm-Anwendung"),
        "fig:case-france": (15, (0.16, 0.23, 0.84, 0.67), 10.3, "Annotierte gerichtete Flussansicht mit Frankreich-Auswahl"),
        "fig:case-price": (17, (0.08, 0.075, 0.92, 0.515), 15.5, "Annotierte Zeitreihe am monatlichen Preisminimum"),
        "fig:case-matrix": (17, (0.08, 0.60, 0.92, 0.86), 15.5, "Annotierte Pixelmatrix mit Dänemark- und Stundenauswahl"),
    }
    page, box, width, alt = crops[label]
    target = temp_dir / f"{label.replace(':', '_')}.png"
    crop_report_page(page, box, target)
    add_figure(doc, target, FIGURE_CAPTIONS[label], width, alt)


def parse_list(lines):
    items = []
    current = []
    for line in lines:
        stripped = line.strip()
        if stripped.startswith("\\item"):
            if current:
                items.append(clean_tex(" ".join(current)))
            current = [stripped[len("\\item"):].strip()]
        elif stripped and not stripped.startswith("\\begin") and not stripped.startswith("\\end"):
            current.append(stripped)
    if current:
        items.append(clean_tex(" ".join(current)))
    return items


def parse_description(lines):
    items = []
    label = None
    body = []
    for line in lines:
        stripped = line.strip()
        match = re.match(r"\\item\[([^\]]+)\]\s*(.*)", stripped)
        if match:
            if label is not None:
                items.append((clean_tex(label), clean_tex(" ".join(body))))
            label, first = match.groups()
            body = [first]
        elif label is not None and stripped and not stripped.startswith(("\\begin", "\\end")):
            body.append(stripped)
    if label is not None:
        items.append((clean_tex(label), clean_tex(" ".join(body))))
    return items


def process_section(doc, path, temp_dir, section_numbers):
    lines = path.read_text(encoding="utf-8").splitlines()
    paragraph = []
    subsection_count = 0
    subsub_count = 0
    current_section = None

    def flush():
        nonlocal paragraph
        if paragraph:
            text = " ".join(x.strip() for x in paragraph)
            if text:
                italic = text.startswith("\\emph{") and text.endswith("}")
                add_body_paragraph(doc, text, italic=italic, centered=italic)
            paragraph = []

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        section_match = re.match(r"\\section\{(.+)\}", stripped)
        subsection_match = re.match(r"\\subsection(\*)?\{(.+)\}", stripped)
        subsub_match = re.match(r"\\subsubsection(\*)?\{(.+)\}", stripped)
        if section_match:
            flush()
            current_section = section_numbers.pop(0)
            subsection_count = 0
            title = f"{current_section} {clean_tex(section_match.group(1))}"
            add_heading(doc, title, 1)
            i += 1
            continue
        if subsection_match:
            flush()
            starred, title = subsection_match.groups()
            if starred:
                visible = clean_tex(title)
            else:
                subsection_count += 1
                subsub_count = 0
                visible = f"{current_section}.{subsection_count} {clean_tex(title)}"
            add_heading(doc, visible, 2)
            i += 1
            continue
        if subsub_match:
            flush()
            starred, title = subsub_match.groups()
            if starred:
                visible = clean_tex(title)
            else:
                subsub_count += 1
                visible = f"{current_section}.{subsection_count}.{subsub_count} {clean_tex(title)}"
            add_heading(doc, visible, 3)
            i += 1
            continue
        if stripped.startswith("\\label"):
            i += 1
            continue
        env_match = re.match(r"\\begin\{(figure|table|longtable|lstlisting|itemize|enumerate|description|quote)\}", stripped)
        if env_match:
            flush()
            env = env_match.group(1)
            block = [line]
            i += 1
            while i < len(lines) and not lines[i].strip().startswith(f"\\end{{{env}}}"):
                block.append(lines[i])
                i += 1
            if i < len(lines):
                block.append(lines[i])
            block_text = "\n".join(block)
            if env in ("table", "longtable"):
                labels = re.findall(r"\\label\{([^{}]+)\}", block_text)
                label = labels[0] if labels else ("app:feedback" if env == "longtable" else "")
                insert_special_table(doc, label)
            elif env == "figure":
                labels = re.findall(r"\\label\{([^{}]+)\}", block_text)
                if labels:
                    insert_special_figure(doc, labels[0], temp_dir)
            elif env == "lstlisting":
                raw = "\n".join(block[1:-1])
                add_code_block(doc, raw)
            elif env == "description":
                for label, body in parse_description(block):
                    p = doc.add_paragraph()
                    p.paragraph_format.left_indent = Cm(0.35)
                    p.paragraph_format.first_line_indent = Cm(-0.35)
                    p.paragraph_format.keep_together = True
                    lead = p.add_run(label + " ")
                    set_font(lead, size=10.5, bold=True)
                    run = p.add_run(body)
                    set_font(run, size=10.5)
            elif env == "quote":
                quote_text = " ".join(
                    x.strip() for x in block[1:-1]
                    if x.strip() and x.strip() != "\\itshape"
                )
                add_body_paragraph(doc, quote_text, italic=True, centered=True)
            else:
                for number, item in enumerate(parse_list(block), start=1):
                    p = doc.add_paragraph()
                    p.paragraph_format.left_indent = Cm(0.65)
                    p.paragraph_format.first_line_indent = Cm(-0.45)
                    p.paragraph_format.space_after = Pt(3)
                    prefix = "• " if env == "itemize" else f"{number}. "
                    lead = p.add_run(prefix)
                    set_font(lead, size=10.5, bold=env == "enumerate")
                    run = p.add_run(item)
                    set_font(run, size=10.5)
            i += 1
            continue
        if stripped.startswith(("\\clearpage", "\\newpage", "\\appendix", "\\toprule", "\\midrule", "\\bottomrule")):
            flush()
            i += 1
            continue
        if not stripped:
            flush()
        else:
            paragraph.append(line)
        i += 1
    flush()


def add_title_page(doc):
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Deutschlands Rolle im europäischen Stromnetz")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(42)
    run = subtitle.add_run("Eine interaktive Visual-Analytics-Anwendung mit Elm")
    set_font(run, size=15)
    for label, value in (("Autor", "Jann Körner"), ("Modul", "Informationsvisualisierung"), ("Semester", "Sommersemester 2026"), ("Abgabe", "21. September 2026")):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        a = p.add_run(f"{label}: ")
        set_font(a, size=11, bold=True)
        b = p.add_run(value)
        set_font(b, size=11)
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_hyperlink(p, "Quellcode und Anwendung im Hochschul-GitLab", "https://gitlab-lehre.informatik.uni-halle.de/aquhr/energiechart-visualisierungsprojekt")
    doc.add_page_break()


def add_front_matter(doc):
    p = doc.add_paragraph()
    run = p.add_run("Inhaltsverzeichnis")
    set_font(run, size=16, bold=True)
    entries = [
        ("1 Einleitung", 4, 0),
        ("1.1 Anwendungshintergrund", 4, 1),
        ("1.2 Zielgruppen", 5, 1),
        ("1.3 Überblick und Beiträge", 5, 1),
        ("2 Daten", 5, 0),
        ("Datenbestand und Einheiten", 5, 1),
        ("Auswahl des Zeitraums", 6, 1),
        ("Abfrage und Vorverarbeitung", 6, 1),
        ("Normalisiertes Format und Bereitstellung", 7, 1),
        ("Datenqualität und Eignung", 7, 1),
        ("3 Visualisierungen", 7, 0),
        ("3.1 Analyse der Anwendungsaufgaben", 8, 1),
        ("3.2 Anforderungen an die Visualisierungen", 9, 1),
        ("3.3 Präsentation der Visualisierungen", 9, 1),
        ("3.3.1 Visualisierung Eins gerichteter radialer Netzwerkgraph", 9, 2),
        ("3.3.2 Visualisierung Zwei gestapelte Zeitreihe mit Flussdetail", 10, 2),
        ("3.3.3 Visualisierung Drei Pixelmatrix", 10, 2),
        ("3.4 Interaktion", 11, 1),
        ("4 Implementierung", 11, 0),
        ("Gesamtaufbau", 11, 1),
        ("Wichtigste Elm-Datenstrukturen", 11, 1),
        ("SVG-Berechnung", 12, 1),
        ("Implementierungsaufwand und Entscheidungen", 12, 1),
        ("Prüfung und Bereitstellung", 12, 1),
        ("Einsatz von KI", 13, 1),
        ("5 Anwendungsfälle", 13, 0),
        ("5.1 Anwendung Visualisierung Eins", 13, 1),
        ("5.2 Anwendung Visualisierung Zwei", 14, 1),
        ("5.3 Anwendung Visualisierung Drei", 15, 1),
        ("6 Verwandte Arbeiten", 16, 0),
        ("7 Zusammenfassung und Ausblick", 17, 0),
        ("A Git-Historie", 17, 0),
        ("B KI-gestützter Arbeitsprozess", 18, 0),
        ("C Umsetzung des Feedbacks zum zweiten Zwischenstand", 19, 0),
        ("Literatur", 20, 0),
    ]
    for title, page, level in entries:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.45 * level)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.tab_stops.add_tab_stop(Cm(15.3), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        text = p.add_run(title)
        set_font(text, size=8.5 if level < 2 else 8.0, bold=level == 0)
        number = p.add_run(f"\t{page}")
        set_font(number, size=8.5 if level < 2 else 8.0, bold=level == 0)
    doc.add_page_break()
    p = doc.add_paragraph()
    run = p.add_run("Abbildungsverzeichnis")
    set_font(run, size=16, bold=True)
    for caption in FIGURE_CAPTIONS.values():
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.4)
        p.paragraph_format.first_line_indent = Cm(-0.4)
        run = p.add_run(caption)
        set_font(run, size=10)
    doc.add_page_break()


def add_literature(doc):
    add_heading(doc, "Literatur", 1)
    for reference in REFERENCES:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.first_line_indent = Cm(-0.75)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(reference)
        set_font(run, size=9.3)


def build():
    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "Deutschlands Rolle im europäischen Stromnetz"
    doc.core_properties.subject = "Projektbericht Informationsvisualisierung"
    doc.core_properties.author = "Jann Körner"
    doc.core_properties.last_modified_by = "Jann Körner"
    doc.core_properties.comments = ""
    doc.core_properties.created = datetime(2026, 9, 11)
    doc.core_properties.modified = datetime(2026, 9, 11)

    add_title_page(doc)
    add_front_matter(doc)
    files = [
        "01_einleitung.tex",
        "02_daten.tex",
        "03_visualisierungen.tex",
        "04_implementierung.tex",
        "05_anwendungsfaelle.tex",
        "06_verwandte_arbeiten.tex",
        "07_zusammenfassung.tex",
        "anhang_git_ki.tex",
    ]
    section_numbers = ["1", "2", "3", "4", "5", "6", "7", "A", "B", "C"]
    with tempfile.TemporaryDirectory(prefix="report_docx_") as temp:
        temp_dir = Path(temp)
        for filename in files:
            process_section(doc, SOURCE / "sections" / filename, temp_dir, section_numbers)
    add_literature(doc)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
