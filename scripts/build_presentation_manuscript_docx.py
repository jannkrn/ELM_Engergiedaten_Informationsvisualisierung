from __future__ import annotations

import shutil
import tempfile
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
REFERENCE = (
    Path.home()
    / ".codex/plugins/cache/openai-curated-remote/openai-templates/0.1.1/skills"
    / "artifact-template-system-design/assets/reference.docx"
)
OUTPUT = ROOT / "report/presentation/Vortragsmanuskript_Jann_Koerner.docx"
FIGURES = ROOT / "experiments/figures/final"
REPORT_PDF = ROOT / "report/final/bericht.pdf"

NAVY = "082A4A"
MUTED = "5B7085"
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F3F6F9"
GRID = "D9D9D9"
BLACK = "000000"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=GRID, size="6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_row_repeat(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_font(run, name="Helvetica Neue", size=9.5, bold=None, italic=None, color=BLACK) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def style_document(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Helvetica Neue"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Helvetica Neue")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Helvetica Neue")
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    title = doc.styles["Title"]
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title.font.size = Pt(24)
    title.font.bold = True
    title.paragraph_format.space_after = Pt(8)

    for style_name, size in (("Heading 1", 14), ("Heading 2", 11), ("Heading 3", 9.5)):
        style = doc.styles[style_name]
        style.font.name = "Helvetica Neue"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Helvetica Neue")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Helvetica Neue")
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)


def clear_body(doc: Document) -> None:
    body = doc._element.body
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)


def clear_paragraph(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def configure_footer(doc: Document) -> None:
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    clear_paragraph(section.first_page_footer.paragraphs[0])
    footer = section.footer.paragraphs[0]
    clear_paragraph(footer)
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Jann Körner | Vortragsmanuskript | Informationsvisualisierung")
    set_font(run, size=7.5, color=MUTED)


def add_heading(doc: Document, text: str, level=1, page_break=False):
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.page_break_before = page_break
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_body(doc: Document, text: str, bold_start: str | None = None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.keep_together = False
    if bold_start and text.startswith(bold_start):
        first = paragraph.add_run(bold_start)
        set_font(first, bold=True)
        rest = paragraph.add_run(text[len(bold_start) :])
        set_font(rest)
    else:
        run = paragraph.add_run(text)
        set_font(run)
    return paragraph


def add_cue(doc: Document, text: str):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_with_next = True
    label = paragraph.add_run("Zeigen  ")
    set_font(label, size=8.7, bold=True, color=MUTED)
    run = paragraph.add_run(text)
    set_font(run, size=8.7, italic=True, color=MUTED)
    return paragraph


def add_bullet(doc: Document, text: str, level=0):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.2 + (0.2 * level))
    paragraph.paragraph_format.first_line_indent = Inches(-0.15)
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(f"• {text}")
    set_font(run)
    return paragraph


def add_number(doc: Document, number: int, text: str):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.22)
    paragraph.paragraph_format.first_line_indent = Inches(-0.18)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(f"{number}. {text}")
    set_font(run)
    return paragraph


def add_code(doc: Document, code: str):
    lines = code.strip("\n").splitlines()
    for index, line in enumerate(lines):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.18)
        paragraph.paragraph_format.right_indent = Inches(0.18)
        paragraph.paragraph_format.space_before = Pt(4 if index == 0 else 0)
        paragraph.paragraph_format.space_after = Pt(4 if index == len(lines) - 1 else 0)
        paragraph.paragraph_format.keep_together = True
        p_pr = paragraph._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), PALE_GRAY)
        p_pr.append(shd)
        run = paragraph.add_run(line or " ")
        set_font(run, name="Consolas", size=8.2, color=BLACK)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    set_row_repeat(header)
    for index, value in enumerate(headers):
        cell = header.cells[index]
        cell.width = Inches(widths[index])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(value)
        set_font(run, size=8.5, bold=True, color=WHITE)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for column, value in enumerate(values):
            cell = cells[column]
            cell.width = Inches(widths[column])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_shading(cell, PALE_BLUE if row_index % 2 else WHITE)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(value)
            set_font(run, size=8.3)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(1)
    return table


def add_figure(doc: Document, path: Path, width: float, caption: str, alt_text: str):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    inline = run.add_picture(str(path), width=Inches(width))
    doc_pr = inline._inline.docPr
    doc_pr.set("descr", alt_text)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(7)
    cap.paragraph_format.keep_together = True
    cap_run = cap.add_run(caption)
    set_font(cap_run, size=8.0, italic=True, color=MUTED)


def build_architecture_image(target: Path) -> bool:
    try:
        import pypdfium2 as pdfium
    except ImportError:
        return False
    pdf = pdfium.PdfDocument(str(REPORT_PDF))
    page = pdf[11]
    image = page.render(scale=2.1).to_pil().convert("RGB")
    width, height = image.size
    crop = image.crop((int(width * 0.095), int(height * 0.272), int(width * 0.90), int(height * 0.535)))
    crop.save(target)
    page.close()
    pdf.close()
    return True


def add_cover(doc: Document) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(58)
    run = paragraph.add_run("INFORMATIONSVISUALISIERUNG")
    set_font(run, size=11, bold=True, color=MUTED)

    title = doc.add_paragraph(style="Title")
    title.add_run("Vortragsmanuskript")
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(26)
    run = subtitle.add_run("Deutschlands Rolle im europäischen Stromnetz")
    set_font(run, size=17, bold=True, color=BLACK)

    add_table(
        doc,
        ["DAUER", "SPRECHER", "STAND"],
        [["15 bis 20 Minuten", "Jann Körner", "11. September 2026"]],
        [2.2, 2.2, 2.2],
    )

    doc.add_paragraph()
    add_table(
        doc,
        ["Angabe", "Inhalt"],
        [
            ["Projekt", "Interaktive Visual-Analytics-Anwendung mit Elm"],
            ["Forschungsfrage", "Wann ist Deutschland Stromimporteur oder Stromexporteur und wie hängen Flüsse, Erzeugungsmix und Preise zusammen"],
            ["Datengrundlage", "Energy-Charts-Daten für Mai 2025 mit 744 Stunden und elf Partnerländern"],
            ["Zweck", "Sprechtext, Vorführhinweise, Codeerklärung und Antworten auf Rückfragen"],
        ],
        [1.45, 5.15],
    )

    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(22)
    run = paragraph.add_run("Kernaussage")
    set_font(run, size=9.5, bold=True, color=MUTED)
    add_body(
        doc,
        "Die Anwendung verbindet Netzwerkgraph, Zeitreihe und Pixelmatrix über einen gemeinsamen Elm-Zustand. Dadurch führt eine Auswahl direkt von einem Überblick zu den zugehörigen absoluten Werten, ohne physische Flüsse, Handel und Preis gleichzusetzen.",
    )


def add_timing(doc: Document) -> None:
    add_heading(doc, "Überblick und Zeitplan", page_break=True)
    add_body(
        doc,
        "Die normalen Absätze sind als Sprechtext formuliert. Kursiv gesetzte Hinweise beginnen mit Zeigen und beschreiben, was währenddessen auf einer Folie oder in der Anwendung sichtbar sein sollte. Codeblöcke müssen nicht vollständig vorgelesen werden; im Vortrag genügt die Erklärung direkt darunter.",
    )
    add_table(
        doc,
        ["Zeit", "Abschnitt", "Ziel"],
        [
            ["0:00 bis 1:30", "Einstieg und Leitfrage", "Problem und Forschungsfrage verständlich machen"],
            ["1:30 bis 3:30", "Daten und Einheiten", "Quelle, Views, Zeitraum und Vorzeichen erklären"],
            ["3:30 bis 5:00", "Aufgaben und Entwurf", "Auswahl der drei Darstellungen begründen"],
            ["5:00 bis 9:30", "Drei Visualisierungen", "Kodierungen und Nutzen zeigen"],
            ["9:30 bis 11:00", "Interaktion", "Kopplung und Bedienung demonstrieren"],
            ["11:00 bis 14:30", "Elm-Implementierung", "Architektur und zentrale Codeideen erklären"],
            ["14:30 bis 17:30", "Fallanalysen", "Konkrete Ergebnisse mit Zahlen zeigen"],
            ["17:30 bis 19:00", "Grenzen und Fazit", "Aussagekraft einordnen und abschließen"],
        ],
        [1.15, 2.05, 3.4],
    )
    add_heading(doc, "Wenn weniger Zeit zur Verfügung steht", level=2)
    add_body(
        doc,
        "Bei ungefähr 12 Minuten kürze ich die Codeerklärung auf State, Msg und SelectCell. Die Detailformeln für Linienbreite und Matrixfarbe sowie die Einordnung verwandter Arbeiten lasse ich dann weg. Forschungsfrage, Daten, alle drei Ansichten, ein Interaktionsbeispiel, die drei wichtigsten Ergebnisse und die Grenzen bleiben erhalten.",
    )


def add_opening(doc: Document) -> None:
    add_heading(doc, "1 Einstieg und Forschungsfrage", page_break=True)
    add_cue(doc, "Titelfolie mit Projekttitel und anschließend die Gesamtansicht der Anwendung")
    add_body(
        doc,
        "Guten Tag, ich stelle heute mein Visualisierungsprojekt Deutschlands Rolle im europäischen Stromnetz vor. Die Anwendung ist vollständig in Elm umgesetzt und verbindet drei interaktive Darstellungen. Untersucht werden physische Stromflüsse zwischen Deutschland und elf Partnerländern, der deutsche Erzeugungsmix und der Day-Ahead-Strompreis.",
    )
    add_body(
        doc,
        "Der Ausgangspunkt ist, dass Deutschland in jeder Stunde eine andere Rolle einnehmen kann. Gegenüber einem Land kann Deutschland Strom importieren, gegenüber einem anderen gleichzeitig exportieren. Zusätzlich können sich die Richtung des Handels und die Richtung des gemessenen physischen Flusses unterscheiden. Eine einzelne Kennzahl reicht deshalb nicht aus.",
    )
    add_body(
        doc,
        "Meine Forschungsfrage lautet: Wann ist Deutschland Stromimporteur oder Stromexporteur, mit welchen Ländern findet der physische Austausch statt und wie hängen diese Flüsse mit dem Erzeugungsmix und den Preisen zusammen?",
        bold_start="Meine Forschungsfrage lautet: ",
    )
    add_body(
        doc,
        "Das Ziel ist kein automatisches Erklärungsmodell. Die Anwendung soll auffällige Stunden und Zeiträume sichtbar machen und anschließend die absoluten Werte liefern, mit denen eine Vermutung überprüft werden kann.",
    )
    add_figure(
        doc,
        FIGURES / "01_uebersicht.png",
        6.65,
        "Gesamtansicht mit Netzwerkgraph, Zeitreihe und Pixelmatrix",
        "Gesamtansicht der Elm-Anwendung mit drei interaktiv verbundenen Visualisierungen",
    )


def add_data(doc: Document) -> None:
    add_heading(doc, "2 Daten und Einheiten", page_break=True)
    add_cue(doc, "Tabelle mit den vier verwendeten PostgreSQL-Views")
    add_body(
        doc,
        "Als Datengrundlage verwende ich Energy-Charts-Daten. Die Veranstaltung stellt Kopien der API-Daten als PostgreSQL-Views im Schema energycharts bereit. Das Exportskript fragt diese Views über PostgREST ab, begrenzt jede Anfrage und paginiert größere Ergebnismengen.",
    )
    add_table(
        doc,
        ["View", "Bedeutung", "Einheit"],
        [
            ["v_cbpf", "Bilateraler physischer Grenzfluss", "GW"],
            ["v_cbet", "Bilateraler grenzüberschreitender Handel", "GW"],
            ["v_price", "Day-Ahead-Preis der Gebotszone Deutschland und Luxemburg", "EUR/MWh"],
            ["v_totalpower", "Deutsche Nettoerzeugung nach Technologien", "MW, anschließend GW"],
        ],
        [1.15, 4.35, 1.2],
    )
    add_body(
        doc,
        "GW bedeutet Gigawatt und beschreibt Leistung. EUR/MWh bedeutet Euro pro Megawattstunde und beschreibt den Preis einer Energiemenge. Im Vortrag formuliere ich zum Beispiel: 18,3 GW erneuerbare Leistung, das entspricht 50 Prozent der gesamten Erzeugungsleistung. Der Punkt in der Bedienoberfläche trennt die beiden Angaben; er bedeutet keine Multiplikation.",
    )
    add_body(
        doc,
        "Für CBPF und CBET gilt dieselbe Vorzeichenkonvention: Positive Werte bedeuten Import nach Deutschland, negative Werte Export aus Deutschland. v_totalpower ist am offiziellen Energy-Charts-Endpunkt in MW dokumentiert und wird durch Division durch 1000 in GW umgerechnet. Die anderen drei Größen werden ohne Einheitenumrechnung übernommen.",
    )
    add_heading(doc, "Zeitraum und Umfang", level=2)
    add_body(
        doc,
        "Der Datensatz umfasst den vollständigen Mai 2025 in UTC. Das sind 744 Stunden. Für jede Stunde liegen Erzeugungsmix, Preis und die Werte von elf Partnerländern vor. Die Pixelmatrix enthält damit 744 mal 11, also 8184 physische Flusswerte. Der Monat ist lang genug für Tages- und Wochenmuster und bleibt gleichzeitig in einer Gesamtansicht handhabbar.",
    )
    add_body(
        doc,
        "Die Rohdaten sind je nach View viertelstündlich oder stündlich. Werte innerhalb derselben Stunde werden gemittelt, weil Leistung und Preis dargestellt werden. Eine Summe der vier Viertelstundenleistungen würde die Einheit verfälschen. Die Erzeugungstechnologien werden zu Erneuerbaren, Kohle, Gas und Sonstigen zusammengefasst.",
    )


def add_pipeline(doc: Document) -> None:
    add_heading(doc, "3 Datenweg und sichere Bereitstellung", page_break=True)
    add_cue(doc, "Kurzer Datenfluss von PostgREST über das Exportskript zur JSON-Datei und zu Elm")
    add_body(
        doc,
        "Die Browseranwendung meldet sich nicht direkt an der Datenbank an. Passwort oder Token wären sonst im ausgelieferten JavaScript sichtbar. Stattdessen ruft ein lokales Python-Skript die geschützten Daten ab, prüft und normalisiert sie und schreibt public/data/energy.json. Diese Datei wird zusammen mit der statischen Anwendung veröffentlicht.",
    )
    add_body(
        doc,
        "Elm lädt nur die fertige JSON-Datei per HTTP. Damit bleiben Zugangsdaten außerhalb des Frontends und die Visualisierung funktioniert später ohne Datenbankverbindung. Python erzeugt keine Visualisierung; es übernimmt ausschließlich Datenexport und Validierung. Die sichtbaren Diagramme werden als SVG in Elm berechnet.",
    )
    add_code(
        doc,
        """loadDataset : (Result Http.Error Dataset -> msg) -> Cmd msg
loadDataset toMsg =
    Http.get
        { url = "data/energy.json"
        , expect = Http.expectJson toMsg datasetDecoder
        }""",
    )
    add_body(
        doc,
        "Dieser Ausschnitt aus Api.elm zeigt den öffentlichen Datenpfad. Http.get lädt energy.json. Http.expectJson verbindet die Antwort mit dem JSON-Decoder und verpackt Erfolg oder Fehler in eine Elm-Nachricht. Zugangsdaten stehen an dieser Stelle bewusst nicht im Code.",
    )
    add_heading(doc, "Prüfungen", level=2)
    add_body(
        doc,
        "Das Validierungsskript kontrolliert 744 lückenlose Stunden, identische Partnerabdeckung, endliche Zahlen, nichtnegative Erzeugungsgruppen und grobe Plausibilitätsgrenzen. Danach wird die Elm-Anwendung mit elm make im optimierten Modus kompiliert. Zusätzlich habe ich alle Auswahlwege und Zeitfenster im Browser geprüft.",
    )


def add_design(doc: Document) -> None:
    add_heading(doc, "4 Aufgaben und Auswahl der Visualisierungen", page_break=True)
    add_cue(doc, "Drei Aufgaben A1 bis A3 und danach die Gesamtansicht")
    add_body(
        doc,
        "Die Diagrammtypen wurden nicht zuerst ausgewählt. Zunächst habe ich drei Analyseaufgaben formuliert. A1 fragt für eine Stunde nach Richtung und Größenordnung der Länderflüsse. A2 untersucht zeitliche Muster in Erzeugungsmix, Preis und dem Fluss eines ausgewählten Landes. A3 sucht in vielen Länder-Stunden-Kombinationen nach auffälligen Perioden und Richtungswechseln.",
    )
    add_table(
        doc,
        ["Aufgabe", "Benötigte Sicht", "Gewählte Technik"],
        [
            ["A1", "Gerichtete Beziehungen Deutschlands zu elf Ländern", "Radialer Netzwerkgraph"],
            ["A2", "Mehrere Größen auf einer gemeinsamen Zeitachse", "Gestapelte Zeitreihe mit Flussdetail"],
            ["A3", "Dichter Überblick über 8184 Länder-Stunden-Werte", "Pixelmatrix"],
        ],
        [0.75, 3.7, 2.2],
    )
    add_body(
        doc,
        "Die drei Ansichten erfüllen damit unterschiedliche Aufgaben. Entscheidend ist ihre Kopplung: Ein gefundenes Matrixpixel soll dasselbe Land im Netzwerk und dieselbe Stunde in der Zeitreihe aktivieren. Genau dafür besitzt die Elm-Anwendung einen gemeinsamen Auswahlzustand.",
    )


def add_network(doc: Document) -> None:
    add_heading(doc, "5 Gerichteter Netzwerkgraph", page_break=True)
    add_cue(doc, "Frankreich im Netzwerk auswählen und den Tooltip öffnen")
    add_figure(
        doc,
        FIGURES / "02_frankreich_netzwerk.png",
        4.8,
        "Gerichteter Netzwerkgraph mit ausgewählter Beziehung zu Frankreich",
        "Radialer Netzwerkgraph mit Deutschland im Zentrum und Frankreich als Auswahl",
    )
    add_body(
        doc,
        "Deutschland steht im Zentrum, die elf Partnerländer liegen radial darum. Eine rote Kante mit Pfeil nach Deutschland bedeutet Import. Eine blaue Kante mit Pfeil nach außen bedeutet Export. Die Linienbreite kodiert den absoluten physischen Flussbetrag. Der Tooltip zeigt zusätzlich den exakten CBPF-Wert und den Handelswert aus CBET.",
    )
    add_body(
        doc,
        "Eine geografische Karte wäre zwar vertraut, würde bei nur einem zentralen Land aber viel Fläche für Geografie verbrauchen. Die tatsächliche Entfernung ist außerdem keine Datendimension. Der radiale Graph priorisiert deshalb die Beziehungen; exakte Werte bleiben über den Tooltip erreichbar.",
    )
    add_code(
        doc,
        """isImport =
    flow.value >= 0

color =
    if isImport then "#c44545" else "#326db6"

strokeWidth =
    String.fromFloat (1.5 + min 20 (abs flow.value * 3.2))""",
    )
    add_body(
        doc,
        "Der Code verwendet zuerst das Vorzeichen für die Richtung und Farbe. Für die Breite wird der Betrag genommen, weil minus 3 GW nicht dünner als plus 3 GW sein soll. min 20 begrenzt den zusätzlichen Anteil; zusammen mit 1,5 ergibt sich eine maximale Breite von 21,5 SVG-Einheiten.",
    )


def add_timeseries(doc: Document) -> None:
    add_heading(doc, "6 Zeitreihe mit Erzeugungsmix und Flussdetail", page_break=True)
    add_cue(doc, "Sieben-Tage-Ansicht vom 8. bis 14. Mai mit Frankreich und 11. Mai 11 Uhr")
    add_figure(
        doc,
        FIGURES / "03_negativpreis_zeitreihe.png",
        6.7,
        "Zeitreihe am monatlichen Preisminimum mit gekoppelter Frankreich-Auswahl",
        "Gestapelte Erzeugungszeitreihe mit separatem physischen Frankreich-Fluss und Detailkarten",
    )
    add_body(
        doc,
        "Die obere Darstellung zeigt absolute Erzeugungsleistung in GW. Grün steht für Erneuerbare, Dunkelgrau für Kohle, Orange für Gas und Hellgrau für Sonstige. Die Gesamthöhe aller Flächen ist die gesamte dargestellte Erzeugungsleistung. Die Prozentwerte erscheinen zusätzlich in den Detailfeldern für die ausgewählte Stunde.",
    )
    add_body(
        doc,
        "Der Länderfluss liegt auf einer eigenen symmetrischen GW-Skala unterhalb der Erzeugung. Dadurch bleiben Werte von wenigen GW lesbar, obwohl die deutsche Erzeugung mehrere Zehn GW beträgt. Der Strompreis steht in einem Detailfeld und nicht auf derselben Y-Achse, weil EUR/MWh nicht mit GW vermischt werden darf. Die gestrichelte Vertikale verbindet alle Werte derselben Stunde.",
    )
    add_code(
        doc,
        """renewTop sample = sample.generation.renewables
coalTop sample = renewTop sample + sample.generation.coal
gasTop sample = coalTop sample + sample.generation.gas
allTop sample = gasTop sample + sample.generation.other

area upper lower color =
    path [ A.d (areaPath x yGeneration samples upper lower)
         , A.fill color ] []""",
    )
    add_body(
        doc,
        "Für jede gestapelte Fläche berechnet Elm eine obere und eine untere Grenze. areaPath verbindet die oberen Punkte von links nach rechts und die unteren Punkte in umgekehrter Reihenfolge. Mit Z wird der SVG-Pfad geschlossen. So entstehen echte Flächen und keine vorgefertigte Diagrammbibliothek.",
    )


def add_matrix(doc: Document) -> None:
    add_heading(doc, "7 Pixelmatrix", page_break=True)
    add_cue(doc, "48-Stunden-Ansicht vom 3. bis 4. Mai und Zelle Dänemark am 3. Mai 12 Uhr")
    add_figure(
        doc,
        FIGURES / "04_daenemark_matrix.png",
        6.7,
        "Pixelmatrix mit markierter Dänemark-Zeile und ausgewählter Stunde",
        "Pixelmatrix der physischen Stromflüsse mit ausgewählter Dänemark-Zelle",
    )
    add_body(
        doc,
        "Jede Zeile steht für ein Partnerland und jede Spalte für eine Stunde. Blau bedeutet Export aus Deutschland, Rot Import nach Deutschland und Hellgrau einen Wert nahe null. Alle Zellen des gesamten Monats verwenden dieselbe globale Grenze. Ein Farbton ist deshalb über Länder und Zeitfenster hinweg vergleichbar.",
    )
    add_body(
        doc,
        "Ein Klick auf eine Zelle wählt gleichzeitig Land und Stunde. Die goldene Zeile und Spalte bilden ein Fadenkreuz; die aktive Zelle erhält zusätzlich eine dunkle Kontur. Die Matrix ist für Muster und Ausnahmen gedacht. Den genauen Wert liest man anschließend im Tooltip oder in den gekoppelten Detailansichten.",
    )
    add_code(
        doc,
        r"""maxAbs =
    countries
        |> List.concatMap (\country ->
            List.map (flowFor country >> abs) scaleSamples)
        |> List.maximum
        |> Maybe.withDefault 1

ratio = min 1 (abs value / maxAbs)""",
    )
    add_body(
        doc,
        "scaleSamples enthält den ganzen Monat, auch wenn nur 48 Stunden sichtbar sind. maxAbs ist deshalb für jede Ansicht gleich. ratio normiert den Betrag zwischen null und eins. Danach mischt colorFor Hellgrau mit Rot oder Blau. Dadurch verändert ein Wechsel des Zeitfensters nicht heimlich die Bedeutung der Farben.",
    )


def add_interaction(doc: Document) -> None:
    add_heading(doc, "8 Kopplung und Live Demonstration", page_break=True)
    add_cue(doc, "Anwendung im Browser bedienen")
    add_body(
        doc,
        "Für die Vorführung verwende ich drei unabhängige Beispiele. Ich stelle klar, dass sie keine zwingende Klickfolge bilden. Zuerst zeige ich die Länderwahl im Netzwerk. Danach wähle ich für die Preisbetrachtung ein eigenes Wochenfenster. Abschließend beginne ich mit einem neuen 48-Stunden-Fenster und wähle eine Matrixzelle.",
    )
    add_number(doc, 1, "Netzwerk: 1. Mai 2025 um 00 Uhr anzeigen und Frankreich anklicken. Die anderen Kanten werden abgeblendet, Frankreich wird hervorgehoben und die violette Flusslinie erscheint.")
    add_number(doc, 2, "Zeitreihe: Auf sieben Tage wechseln, den Zeitraum 8. bis 14. Mai anzeigen, Frankreich wählen und den 11. Mai um 11 Uhr anklicken. Die gemeinsame Zeitmarke und alle Detailwerte wechseln.")
    add_number(doc, 3, "Matrix: Auf 48 Stunden wechseln, zum 3. bis 4. Mai navigieren und Dänemark am 3. Mai um 12 Uhr wählen. Partnerland und Stunde werden gemeinsam gesetzt.")
    add_body(
        doc,
        "Falls während der Vorführung eine falsche Auswahl aktiv ist, nutze ich Auswahl zurücksetzen. Die Zeitfenster 48 Stunden, sieben Tage und gesamter Monat sind bewusst fest definiert. Dadurch bleibt die Matrix lesbar und die Navigation vorhersehbar.",
    )
    add_code(
        doc,
        """( SelectPartner country, Ready state ) ->
    Ready { state | selectedPartner = Just country }

( SelectTime index, Ready state ) ->
    Ready { state | selectedIndex = index }

( SelectCell country index, Ready state ) ->
    Ready
        { state
            | selectedPartner = Just country
            , selectedIndex = index
        }""",
    )
    add_body(
        doc,
        "Diese drei Update-Zweige erklären die Kopplung. Eine Netzwerkauswahl ändert nur selectedPartner, eine Zeitwahl nur selectedIndex. Eine Matrixzelle setzt beide Felder atomar. Alle Ansichten erhalten danach denselben neuen State und können daher nicht mit unterschiedlichen Auswahlzuständen weiterlaufen.",
    )


def add_architecture(doc: Document, architecture_path: Path | None) -> None:
    add_heading(doc, "9 Elm Architektur und Module", page_break=True)
    add_cue(doc, "Architekturdiagramm und danach den State-Code")
    add_body(
        doc,
        "Die Anwendung folgt The Elm Architecture. Das Model enthält den gesamten Zustand. Msg beschreibt mögliche Ereignisse. update erzeugt aus einer Nachricht einen neuen Zustand. view rendert daraus die Oberfläche. Da Elm unveränderliche Daten verwendet, entsteht bei jeder Interaktion ein neuer konsistenter State.",
    )
    if architecture_path and architecture_path.exists():
        add_figure(
            doc,
            architecture_path,
            6.3,
            "Daten- und Zustandsfluss der Elm-Anwendung",
            "Moduldiagramm vom JSON-Datensatz über Api und Main zu drei zustandslosen Ansichten",
        )
    add_table(
        doc,
        ["Modul", "Aufgabe", "Wichtiger Punkt"],
        [
            ["Domain.elm", "Typen für Datensatz, Stundenwerte, Erzeugung und Flüsse", "Gemeinsames fachliches Modell"],
            ["Api.elm", "HTTP-Abruf und JSON-Decoder", "Fehler wird als Result an Main gemeldet"],
            ["Main.elm", "Model, Update, Zeitfenster und Kopplung", "Ein gemeinsamer Auswahlzustand"],
            ["FlowNetwork.elm", "Gerichtete und gewichtete SVG-Kanten", "Landwahl über Callback"],
            ["TimeSeries.elm", "Gestapelte Flächen, Flusslinie und Detailwerte", "Stundenwahl über Callback"],
            ["FlowMatrix.elm", "Globale Farbskala und Zellen", "Land und Stunde gemeinsam wählen"],
        ],
        [1.35, 3.25, 2.1],
    )


def add_elm_state(doc: Document) -> None:
    add_heading(doc, "10 Zentrale Elm Datentypen", page_break=True)
    add_cue(doc, "State und Msg nebeneinander erläutern")
    add_code(
        doc,
        """type Model
    = Loading
    | Failed String
    | Ready State

type alias State =
    { dataset : Dataset
    , selectedIndex : Int
    , selectedPartner : Maybe String
    , windowStart : Int
    , windowSize : Int
    }""",
    )
    add_body(
        doc,
        "Model trennt Laden, Fehler und die fertige Anwendung. Im Ready-State liegt der Datensatz zusammen mit der globalen Stundenposition, dem optionalen Partnerland und dem sichtbaren Zeitfenster. Maybe String ist wichtig: Die Erzeugungszeitreihe funktioniert auch dann, wenn noch kein Land gewählt wurde.",
    )
    add_code(
        doc,
        """type Msg
    = GotDataset (Result Http.Error Dataset)
    | SelectPartner String
    | SelectTime Int
    | SelectCell String Int
    | SetWindowSize Int
    | MoveWindow Int
    | Reset""",
    )
    add_body(
        doc,
        "Jede Benutzeraktion besitzt eine eigene Nachricht. Das macht die Zustandsübergänge nachvollziehbar und testbar. Besonders wichtig ist die Trennung zwischen SelectPartner, SelectTime und SelectCell. Dadurch ist klar, welche Dimension eine Interaktion verändert.",
    )
    add_code(
        doc,
        r"""View.TimeSeries.view visibleSamples visibleSelectedIndex
    state.selectedPartner
    (\localIndex -> SelectTime (state.windowStart + localIndex))

View.FlowMatrix.view countryList samples visibleSamples
    visibleSelectedIndex state.selectedPartner
    (\country localIndex ->
        SelectCell country (state.windowStart + localIndex))""",
    )
    add_body(
        doc,
        "Die Ansichten arbeiten mit einem lokalen Index ihres sichtbaren Fensters. Main addiert windowStart und übersetzt ihn damit in den globalen Index der 744 Stunden. Diese Umrechnung verhindert, dass ein Klick im zweiten Wochenfenster versehentlich eine Stunde aus dem Monatsanfang auswählt.",
    )


def add_findings(doc: Document) -> None:
    add_heading(doc, "11 Fallanalysen und Ergebnisse", page_break=True)
    add_cue(doc, "Jeweils die passende annotierte Abbildung oder den gespeicherten Zustand zeigen")
    add_heading(doc, "Frankreich am 1 Mai um 00 Uhr", level=2)
    add_body(
        doc,
        "Der physische Fluss aus Frankreich beträgt plus 0,058 GW und zeigt damit einen kleinen Import nach Deutschland. Der Handelswert beträgt gleichzeitig minus 2,940 GW und steht für kommerziellen Export. Das Beispiel zeigt, warum physischer Fluss und Handel nicht synonym verwendet werden dürfen.",
    )
    add_heading(doc, "Preisminimum am 11 Mai um 11 Uhr", level=2)
    add_body(
        doc,
        "Der Day-Ahead-Preis erreicht minus 250,32 EUR/MWh. Die erneuerbare Leistung beträgt 58,961 GW und damit 88,7 Prozent der dargestellten Gesamterzeugung von ungefähr 66,5 GW. Gleichzeitig importiert Deutschland physisch plus 1,382 GW aus Frankreich.",
    )
    add_body(
        doc,
        "Über den gesamten Mai korreliert der Preis mit der absoluten erneuerbaren Leistung mit r gleich minus 0,758 und mit dem erneuerbaren Anteil mit r gleich minus 0,817. Die einfache Summe der elf physischen Länderflüsse korreliert mit dem Preis dagegen nur mit r gleich 0,084. Diese Werte beschreiben Zusammenhänge im Mai, aber keine Ursache.",
    )
    add_heading(doc, "Dänemark am 3 Mai um 12 Uhr", level=2)
    add_body(
        doc,
        "Die ausgewählte Matrixzelle zeigt minus 1,710 GW und damit Export aus Deutschland nach Dänemark. In derselben Stunde sind viele andere Partner ebenfalls blau oder nahe null, während Frankreich mit plus 0,582 GW rötlich ist. Deutschland kann also gleichzeitig gegenüber verschiedenen Ländern unterschiedliche Rollen einnehmen.",
    )
    add_heading(doc, "So formuliere ich die Interpretation", level=2)
    add_body(
        doc,
        "Ich spreche von einem sichtbaren oder statistischen Zusammenhang, nicht von einem Kausalnachweis. Nachfrage, Wetter, installierte Leistung, Netzengpässe und Marktmechanismen fehlen im Datensatz. Die Visualisierung hilft, eine Hypothese zu finden und anhand konkreter Werte zu prüfen; sie erklärt nicht allein, warum ein Preis oder Fluss entstanden ist.",
    )


def add_related_and_limits(doc: Document) -> None:
    add_heading(doc, "12 Fachliche Einordnung und Grenzen", page_break=True)
    add_cue(doc, "Vier kurze Literaturbezüge und anschließend die Grenzen")
    add_body(
        doc,
        "Der Entwurf folgt dem verschachtelten Modell von Tamara Munzner: Zuerst wurden Domänenproblem und Aufgaben bestimmt, dann die visuellen Kodierungen und anschließend ihre technische Umsetzung. Aigner und Kollegen begründen die gemeinsame und geordnete Zeitachse. Keim beschreibt die platzsparende Stärke pixelorientierter Verfahren. North und Shneiderman liefern die Grundlage für koordinierte Ansichten und Brushing-and-Linking.",
    )
    add_body(
        doc,
        "Verwandte Visual-Analytics-Arbeiten zu Smart Grids koppeln ebenfalls mehrere Ansichten. Mein Projekt unterscheidet sich durch den retrospektiven Fokus auf Länderflüsse, Erzeugungsmix und Preis. Es enthält kein Alarmmodell und gibt keine Handlungsempfehlungen.",
    )
    add_heading(doc, "Grenzen der Untersuchung", level=2)
    add_bullet(doc, "Der Mai 2025 erlaubt keine saisonalen oder mehrjährigen Aussagen.")
    add_bullet(doc, "v_totalpower umfasst auch industrielle Eigenerzeugung.")
    add_bullet(doc, "Zeitgleich auftretende Erzeugung und Grenzflüsse zeigen nicht die physische Herkunft einzelner Strommengen.")
    add_bullet(doc, "Die Summe der elf bilateralen Flüsse ist eine eigene explorative Kennzahl und kein separates Energy-Charts-Messfeld.")
    add_bullet(doc, "Freies Zoomen, Mehrfachauswahl und Sortierung der Matrixzeilen wurden nicht umgesetzt.")
    add_heading(doc, "Sinnvolle Erweiterungen", level=2)
    add_body(
        doc,
        "Sinnvolle nächste Schritte wären frei wählbare Start- und Enddaten, ein Drill-down auf einzelne Erzeugungstechnologien, ein Scatterplot für Preis und erneuerbare Leistung, eine Differenzansicht für CBPF und CBET sowie saisonale oder mehrjährige Analysen.",
    )


def add_conclusion(doc: Document) -> None:
    add_heading(doc, "13 Fazit", page_break=True)
    add_cue(doc, "Gesamtansicht oder eine Folie mit den drei Ansichten")
    add_body(
        doc,
        "Ich fasse zusammen: Der Netzwerkgraph beantwortet, mit welchen Ländern Deutschland in einer Stunde physisch verbunden ist und in welche Richtung Leistung fließt. Die Zeitreihe zeigt, wie sich Erzeugungsmix, Preis und ein ausgewählter Länderfluss zeitlich entwickeln. Die Pixelmatrix verdichtet den Monat zu einem Überblick und dient gleichzeitig als Auswahlwerkzeug.",
    )
    add_body(
        doc,
        "Der wichtigste technische Punkt ist der gemeinsame Elm-Zustand. Land, Stunde und Zeitfenster werden zentral verwaltet und an alle zustandslosen Ansichten weitergegeben. Dadurch wirkt jede Auswahl über die Grenzen eines einzelnen Diagramms hinaus.",
    )
    add_body(
        doc,
        "Inhaltlich zeigt die Fallanalyse einen deutlichen negativen Zusammenhang zwischen erneuerbarer Erzeugung und Preis im Mai 2025. Sie zeigt außerdem, dass physischer Fluss und Handel unterschiedliche Richtungen haben können und dass Deutschland in derselben Stunde gegenüber verschiedenen Ländern importieren und exportieren kann.",
    )
    add_body(
        doc,
        "Damit endet mein Vortrag. Vielen Dank für Ihre Aufmerksamkeit. Ich beantworte gern Ihre Fragen.",
    )


def add_questions(doc: Document) -> None:
    add_heading(doc, "Fragen und kurze Antworten", page_break=True)
    questions = [
        ("Warum kein geografischer Kartenausschnitt", "Die Entfernung zwischen Ländern ist keine Datendimension. Für elf Beziehungen zu einem Zentrum zeigt der radiale Graph Richtung und Betrag platzsparender. Eine Karte wäre eine sinnvolle spätere Ergänzung für geografische Orientierung."),
        ("Warum wird der Preis nicht als Linie über die Erzeugung gelegt", "Preis wird in EUR/MWh gemessen, Erzeugung in GW. Eine gemeinsame Y-Achse wäre fachlich falsch. Der Preis steht deshalb im Detailfeld derselben ausgewählten Stunde."),
        ("Warum wird Python verwendet, obwohl Elm gefordert ist", "Python übernimmt nur geschützten Datenabruf, Normalisierung und Validierung. Alle drei interaktiven Visualisierungen, ihre SVG-Geometrie und ihre Kopplung sind in Elm implementiert."),
        ("Was ist der Unterschied zwischen CBPF und CBET", "CBPF beschreibt den gemessenen physischen Grenzfluss. CBET beschreibt den kommerziellen grenzüberschreitenden Handel. Beide können gleichzeitig unterschiedliche Vorzeichen besitzen."),
        ("Warum Mai 2025", "Ein vollständiger Monat enthält 744 Stunden, mehrere Tages- und Wochenmuster sowie positive und negative Preise. Er ist groß genug für Exploration und bleibt in der Pixelmatrix handhabbar."),
        ("Sind die Farben der Matrix zwischen Zeilen vergleichbar", "Ja. Das Maximum wird aus allen 8184 Monatswerten gebildet. Es hängt nicht vom sichtbaren Fenster oder vom Partnerland ab."),
        ("Was bedeutet ein negativer Preis", "Der Day-Ahead-Marktpreis liegt unter null. Das kann bei hohem Angebot, geringer Nachfrage oder begrenzter Flexibilität auftreten. Die Anwendung zeigt Zusammenhänge, beweist aber keine einzelne Ursache."),
        ("Wie werden Zugangsdaten geschützt", "Nur das lokale Exportskript verwendet Token oder Passwort. Die veröffentlichte Elm-Anwendung lädt eine vorbereitete JSON-Datei und enthält keine Datenbankzugangsdaten."),
        ("Was würdest du als Nächstes verbessern", "Ich würde frei wählbare Zeiträume, einen Drill-down auf Einzeltechnologien und einen Scatterplot für Preis-Erzeugungs-Zusammenhänge ergänzen."),
    ]
    for question, answer in questions:
        add_heading(doc, question, level=2)
        add_body(doc, answer)

    add_heading(doc, "Quellen für Rückfragen", level=2)
    for source in [
        "Energy-Charts API und OpenAPI-Spezifikation",
        "PostgREST-Dokumentation zu Tabellen, Views, Limits und Pagination",
        "Elm Guide zu The Elm Architecture",
        "Tamara Munzner zu Design und Validierung von Visualisierungen",
        "Aigner, Miksch, Schumann und Tominski zu zeitbezogenen Daten",
        "Daniel A. Keim zu pixelorientierten Visualisierungstechniken",
        "North und Shneiderman zu koordinierten Ansichten",
    ]:
        add_bullet(doc, source)


def build() -> None:
    if not REFERENCE.exists():
        raise FileNotFoundError(f"Template not found: {REFERENCE}")
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(REFERENCE, OUTPUT)
    doc = Document(OUTPUT)
    clear_body(doc)
    style_document(doc)
    configure_footer(doc)
    doc.core_properties.title = "Vortragsmanuskript Deutschlands Rolle im europäischen Stromnetz"
    doc.core_properties.subject = "Projektvortrag Informationsvisualisierung"
    doc.core_properties.author = "Jann Körner"
    doc.core_properties.keywords = "Elm, Visual Analytics, Energy Charts, Stromnetz"

    with tempfile.TemporaryDirectory(prefix="presentation_docx_") as temp_dir:
        architecture_path = Path(temp_dir) / "architecture.png"
        has_architecture = build_architecture_image(architecture_path)

        add_cover(doc)
        add_timing(doc)
        add_opening(doc)
        add_data(doc)
        add_pipeline(doc)
        add_design(doc)
        add_network(doc)
        add_timeseries(doc)
        add_matrix(doc)
        add_interaction(doc)
        add_architecture(doc, architecture_path if has_architecture else None)
        add_elm_state(doc)
        add_findings(doc)
        add_related_and_limits(doc)
        add_conclusion(doc)
        add_questions(doc)

        doc.save(OUTPUT)

    print(OUTPUT)


if __name__ == "__main__":
    build()
